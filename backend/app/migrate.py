"""
Migration script: index documents from a volume directory into Qdrant.

Usage:
    python -m app.migrate                    # uses env vars
    python -m app.migrate --dir /data/docs   # override volume path
    python -m app.migrate --collection mydocs
    python -m app.migrate --recreate         # drop & recreate collection
"""

import os
import sys
import json
import argparse
import hashlib
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance,
    VectorParams,
    PointStruct,
    Filter,
    FieldCondition,
    MatchValue,
)

load_dotenv()

EMBEDDING_MODEL = os.getenv("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small")
EMBEDDING_BASE_URL = os.getenv("OPENAI_EMBEDDING_BASE_URL", "https://api.openai.com/v1")
EMBEDDING_API_KEY = os.getenv("OPENAI_EMBEDDING_API_KEY")
QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
QDRANT_API_KEY = os.getenv("QDRANT_API_KEY", None)
DOCUMENTS_VOLUME = os.getenv("DOCUMENTS_VOLUME", "/data/documents")

VECTOR_SIZE = 512
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
COLLECTION_NAME = "documents"

SUPPORTED_EXTENSIONS = {
    ".txt": "txt",
    ".md": "md",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".csv": "csv",
}


def parse_args():
    parser = argparse.ArgumentParser(description="Index documents into Qdrant")
    parser.add_argument("--dir", default=DOCUMENTS_VOLUME, help="Documents volume path")
    parser.add_argument("--collection", default=COLLECTION_NAME, help="Qdrant collection name")
    parser.add_argument("--recreate", action="store_true", help="Drop and recreate collection")
    parser.add_argument("--force", action="store_true", help="Re-index even if already indexed")
    return parser.parse_args()


def get_embeddings():
    return OpenAIEmbeddings(
        model=EMBEDDING_MODEL,
        base_url=EMBEDDING_BASE_URL,
        api_key=EMBEDDING_API_KEY,
        dimensions=VECTOR_SIZE,
    )


def get_qdrant():
    return QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)


def ensure_collection(client: QdrantClient, name: str, recreate: bool = False):
    if recreate:
        try:
            client.delete_collection(name)
            print(f"  Deleted collection '{name}'")
        except Exception:
            pass

    collections = [c.name for c in client.get_collections().collections]
    if name not in collections:
        client.create_collection(
            collection_name=name,
            vectors_config=VectorParams(size=VECTOR_SIZE, distance=Distance.COSINE),
        )
        print(f"  Created collection '{name}' (size={VECTOR_SIZE})")
    else:
        print(f"  Collection '{name}' already exists")


def get_existing_sources(client: QdrantClient, name: str) -> set:
    """Return set of source_path values already indexed."""
    sources = set()
    next_offset = None
    while True:
        result = client.scroll(
            collection_name=name,
            limit=1000,
            offset=next_offset,
            with_payload=["source_path"],
            with_vectors=False,
        )
        points, next_offset = result
        for p in points:
            if p.payload and p.payload.get("source_path"):
                sources.add(p.payload["source_path"])
        if next_offset is None or next_offset == "" or next_offset == 0:
            break
    return sources


def read_file(filepath: Path) -> str | None:
    ext = filepath.suffix.lower()
    try:
        if ext == ".txt" or ext == ".md" or ext == ".csv":
            return filepath.read_text(encoding="utf-8", errors="replace")
        elif ext == ".pdf":
            return _read_pdf(filepath)
        elif ext == ".docx":
            return _read_docx(filepath)
        elif ext == ".xlsx":
            return _read_xlsx(filepath)
    except Exception as e:
        print(f"    Error reading {filepath}: {e}")
        return None


def _read_pdf(filepath: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(filepath))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages) if pages else ""


def _read_docx(filepath: Path) -> str:
    from docx import Document

    doc = Document(str(filepath))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx(filepath: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(filepath), read_only=True, data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"--- Sheet: {sheet_name} ---")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def get_file_metadata(filepath: Path) -> dict:
    stat = filepath.stat()
    created = datetime.fromtimestamp(stat.st_ctime).isoformat()
    modified = datetime.fromtimestamp(stat.st_mtime).isoformat()
    return {
        "created_at": created,
        "modified_at": modified,
    }


def main():
    args = parse_args()
    docs_dir = Path(args.dir).resolve()
    collection = args.collection

    if not docs_dir.exists():
        print(f"Error: directory '{docs_dir}' does not exist")
        sys.exit(1)

    print(f"Documents directory: {docs_dir}")
    print(f"Qdrant URL: {QDRANT_URL}")
    print(f"Collection: {collection}")

    embeddings = get_embeddings()
    client = get_qdrant()

    ensure_collection(client, collection, recreate=args.recreate)

    existing_sources = set()
    if not args.force:
        existing_sources = get_existing_sources(client, collection)
        print(f"  Already indexed: {len(existing_sources)} source paths")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    all_files = []
    for ext in SUPPORTED_EXTENSIONS:
        all_files.extend(docs_dir.rglob(f"*{ext}"))

    total = len(all_files)
    indexed = 0
    skipped = 0

    print(f"\nFound {total} files. Starting indexing...\n")

    for filepath in sorted(all_files):
        rel_path = str(filepath.relative_to(docs_dir))
        doc_type = SUPPORTED_EXTENSIONS.get(filepath.suffix.lower(), "unknown")

        if not args.force and rel_path in existing_sources:
            skipped += 1
            continue

        print(f"  [{indexed + 1}/{total}] {rel_path}")
        text = read_file(filepath)
        if not text or not text.strip():
            print(f"    -> empty, skipped")
            skipped += 1
            continue

        chunks = splitter.split_text(text)
        meta = get_file_metadata(filepath)
        print(f"    -> {len(chunks)} chunks")

        points = []
        for chunk_idx, chunk in enumerate(chunks):
            point_id = hashlib.md5(f"{rel_path}:{chunk_idx}".encode()).hexdigest()

            points.append(
                PointStruct(
                    id=point_id,
                    vector=[0.0] * VECTOR_SIZE,
                    payload={
                        "content": chunk,
                        "source_path": rel_path,
                        "doc_type": doc_type,
                        "chunk_index": chunk_idx,
                        "created_at": meta["created_at"],
                        "modified_at": meta["modified_at"],
                    },
                )
            )

        batch_size = 20
        for i in range(0, len(points), batch_size):
            batch = points[i : i + batch_size]

            texts = [p.payload["content"] for p in batch]
            vectors = embeddings.embed_documents(texts)

            for p, vec in zip(batch, vectors):
                p.vector = vec

            client.upsert(collection_name=collection, points=batch)

        indexed += len(chunks)

    print(f"\nDone! Indexed {indexed} chunks from {total} files ({skipped} skipped)")


if __name__ == "__main__":
    main()
